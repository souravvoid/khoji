"""PPTX text extraction using python-pptx.

Extracts text from each slide (title + body). Each slide is
represented as a "page" in the ExtractionResult.
"""

from __future__ import annotations

from pathlib import Path

from khoji_engine.pipeline.pdf_extractor import (
    ExtractedPage,
    ExtractionResult,
    _clean_extracted_text,
)


def extract_pptx(file_path: str | Path) -> ExtractionResult:
    """Extract text from a .pptx file, one page per slide."""
    file_path = Path(file_path)
    result = ExtractionResult(filename=file_path.name, page_count=0)

    try:
        from pptx import Presentation
    except ImportError:
        result.errors.append("python-pptx is not installed")
        return result

    try:
        prs = Presentation(str(file_path))
    except Exception as e:
        result.errors.append(f"Failed to open PPTX: {e}")
        return result

    full_text_parts: list[str] = []
    offset = 0

    for slide_idx, slide in enumerate(prs.slides):
        parts: list[str] = []

        if slide.shapes.title and slide.shapes.title.text.strip():
            parts.append(slide.shapes.title.text.strip())

        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    para_text = para.text.strip()
                    if para_text:
                        parts.append(para_text)

        if not parts:
            continue

        slide_text = "\n".join(parts)
        clean = _clean_extracted_text(slide_text)
        if not clean:
            continue

        ep = ExtractedPage(
            page_number=slide_idx + 1,
            text=clean,
            char_offset=offset,
            char_length=len(clean),
        )
        result.pages.append(ep)
        full_text_parts.append(clean)
        offset += len(clean)

    result.page_count = len(result.pages)
    result.full_text = "\n\n".join(full_text_parts)
    return result
